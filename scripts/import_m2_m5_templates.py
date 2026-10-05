from __future__ import annotations

import json
import re
import unicodedata
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
UPLOAD = ROOT.parents[1] / "upload"
OUTPUT = ROOT / "src/constants/m2-m5-recipe-templates.ts"
M1 = ROOT / "src/constants/m1-recipe-templates.ts"


def norm(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def number(value: object) -> float | None:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    try:
        return float(str(value).replace(",", "."))
    except (TypeError, ValueError):
        return None


def category(title: str, source: str) -> str:
    value = norm(f"{source} {title}")
    if any(x in value for x in ("ciorba", "supa", "bors", "crema")): return "soup"
    if "salata" in value: return "salad"
    if any(x in value for x in ("desert", "prajitura", "tarta", "chec", "budinca", "mousse")): return "dessert"
    if any(x in value for x in ("garnitura", "piure", "cartofi", "orez", "mamalig")): return "side"
    if any(x in value for x in ("aperitiv", "starter", "gustare")): return "starter"
    return "main"


existing_titles = {norm(x) for x in re.findall(r'"title":\s*"([^"]+)"', M1.read_text(encoding="utf-8"))}
candidates: dict[str, dict] = {}
files = sorted(UPLOAD.glob("*-M[2345]-*.xlsx"))

for path in files:
    source_book = re.search(r"M[2345]", path.name).group(0)
    workbook = load_workbook(path, read_only=True, data_only=True)
    nutrition: dict[str, tuple[float | None, float | None, int]] = {}
    if "Calorii Ingrediente" in workbook.sheetnames:
        sheet = workbook["Calorii Ingrediente"]
        for row_no, row in enumerate(sheet.iter_rows(values_only=True), 1):
            name = norm(row[0] if row else "")
            if name:
                nutrition[name] = (number(row[1] if len(row) > 1 else None), number(row[3] if len(row) > 3 else None), row_no)

    for sheet in workbook.worksheets:
        if "calorii" in norm(sheet.title) or sheet.max_row < 5: continue
        rows = list(sheet.iter_rows(values_only=True))
        for index, row in enumerate(rows):
            values = list(row)
            headers = [norm(value) for value in values]
            try: name_col = headers.index("denumire produs")
            except ValueError: continue
            quantity_col = next((i for i, value in enumerate(headers) if "cant portie" in value), None)
            if quantity_col is None or index == 0: continue
            title = str(rows[index - 1][name_col] or "").strip()
            if not title or norm(title) in {"denumire produs", "portii"}: continue
            source_category = ""
            if index >= 2:
                source_category = str(rows[index - 2][name_col] or "").strip()
            ingredients = []
            row_no = index + 1
            while row_no < len(rows):
                current = rows[row_no]
                item_name = str(current[name_col] or "").strip() if name_col < len(current) else ""
                item_quantity = number(current[quantity_col] if quantity_col < len(current) else None)
                ordinal = number(current[name_col - 1] if name_col > 0 and name_col - 1 < len(current) else None)
                if not item_name or item_quantity is None or item_quantity <= 0 or ordinal is None: break
                kcal, price, nutrition_row = nutrition.get(norm(item_name), (None, None, None))
                ingredients.append({
                    "name": item_name.lower(),
                    "quantityGrams": round(item_quantity * 1000, 3),
                    "purchasePricePerKg": round(price or 0, 4),
                    "energyKcalPer100g": kcal,
                    "nutritionSourceRow": nutrition_row,
                })
                row_no += 1
            key = norm(title)
            if not key or key in existing_titles or not ingredients: continue
            score = len(ingredients) * 100 + sum(i["energyKcalPer100g"] is not None for i in ingredients) * 10 + sum(i["purchasePricePerKg"] > 0 for i in ingredients)
            energy = round(sum(i["quantityGrams"] * (i["energyKcalPer100g"] or 0) / 100 for i in ingredients), 2)
            item = {"title": title.upper(), "category": category(title, source_category), "sourceCategory": source_category or None, "sourceSheet": f"{source_book}.xlsx · {sheet.title}", "sourceHeaderRow": index + 1, "energyKcalPerPortion": energy, "ingredients": ingredients, "_score": score, "_source": source_book}
            if key not in candidates or score > candidates[key]["_score"]: candidates[key] = item

items = sorted(candidates.values(), key=lambda item: (item["category"], norm(item["title"])))
for index, item in enumerate(items, 1):
    item["id"] = f"m2m5-{index:03d}"
    item.pop("_score", None); item.pop("_source", None)

header = """/**\n * Rețete unice extrase din M2–M5. Generate mecanic; nu edita manual.\n * Dublurile dintre M2–M5 și cele deja existente în M1 au fost eliminate după denumirea normalizată.\n */\nimport type { M1RecipeTemplate } from '@/constants/m1-recipe-templates';\n\nexport const M2_M5_RECIPE_TEMPLATES: readonly M1RecipeTemplate[] = """
OUTPUT.write_text(header + json.dumps(items, ensure_ascii=False, indent=2) + ";\n", encoding="utf-8")
print(json.dumps({"files": [p.name for p in files], "existing_m1": len(existing_titles), "unique_added": len(items)}, ensure_ascii=False))
